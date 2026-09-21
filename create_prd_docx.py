import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, color="D1D5DB", sz="4", val="single"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def add_formatted_runs(paragraph, text):
    # Tokens: `code`, **bold**, *italic*, [text](url)
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))'
    tokens = re.split(pattern, text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(199, 37, 78) # berry code color
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token)
            if m:
                run = paragraph.add_run(m.group(1))
                run.font.color.rgb = RGBColor(27, 94, 32)
                run.underline = True
            else:
                paragraph.add_run(token)
        else:
            paragraph.add_run(token)

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC") # light slate
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="1B5E20"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.strip('\n'))
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F8E9") # light green
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="2E7D32"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    add_formatted_runs(p, text)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_docx():
    doc = docx.Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        
        # Header text
        header_p = section.header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h_run = header_p.add_run("AgriBuddy — Product Requirements Document (PRD) | STSI4440")
        h_run.font.name = 'Calibri'
        h_run.font.size = Pt(8.5)
        h_run.font.color.rgb = RGBColor(120, 120, 120)
        
        # Footer text
        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run = footer_p.add_run("Dokumen Spesifikasi Kebutuhan Perangkat Lunak AgriBuddy • Semester 7")
        f_run.font.name = 'Calibri'
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(120, 120, 120)

    # Styles setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(33, 33, 33)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # --- COVER / TITLE BANNER ---
    banner_table = doc.add_table(rows=1, cols=1)
    banner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner_cell = banner_table.cell(0, 0)
    banner_cell.width = Inches(6.5)
    set_cell_background(banner_cell, "1B5E20") # Deep Forest Green
    set_cell_margins(banner_cell, top=280, bottom=280, left=240, right=240)
    
    p_pre = banner_cell.paragraphs[0]
    p_pre.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_pre.paragraph_format.space_before = Pt(4)
    p_pre.paragraph_format.space_after = Pt(4)
    run_pre = p_pre.add_run("DOKUMEN KEBUTUHAN PRODUK & ARSITEKTUR SISTEM")
    run_pre.font.name = 'Calibri'
    run_pre.font.size = Pt(10)
    run_pre.font.bold = True
    run_pre.font.color.rgb = RGBColor(200, 230, 201) # Soft Light Green
    
    p_title = banner_cell.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("Product Requirements Document (PRD)\nAgriBuddy")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(255, 255, 255)
    
    p_sub = banner_cell.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    run_sub = p_sub.add_run("Ekosistem Pendamping Petani Cerdas 🌾\nArsitektur Web Responsif Terpadu & Ecosystem-First")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(232, 245, 233)
    
    # Metadata info table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Mata Kuliah / Tugas", "STSI4440 — Capstone Project (Semester 7)"),
        ("Repositori Resmi Git", "https://github.com/ombay86/agribuddy"),
        ("Versi Dokumen", "Versi 2.0 (Arsitektur Web & Kolaborasi Ekosistem)"),
        ("Status Implementasi", "100% Selesai & Terverifikasi (Full Stack Integration Passed)")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F1F8E9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=80, bottom=80, left=100, right=100)
        set_cell_margins(c1, top=80, bottom=80, left=100, right=100)
        set_cell_borders(c0, color="C8E6C9")
        set_cell_borders(c1, color="E0E0E0")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(46, 125, 50)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        if "http" in val:
            r1.font.color.rgb = RGBColor(27, 94, 32)
            r1.underline = True
            
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Read markdown content
    with open('PRD_AgriTech_Capstone.md', 'r', encoding='utf-8') as f:
        md_text = f.read()

    lines = md_text.split('\n')
    
    in_code_block = False
    code_block_lines = []
    in_table = False
    table_rows = []
    
    skip_header = True # skip initial title already rendered in banner
    
    for line in lines:
        stripped = line.strip()
        
        # Check title header skip
        if skip_header:
            if stripped.startswith('## 1. Ringkasan Eksekutif'):
                skip_header = False
            else:
                continue

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                add_code_block(doc, '\n'.join(code_block_lines))
                in_code_block = False
                code_block_lines = []
            else:
                in_code_block = True
                code_block_lines = []
            continue

        if in_code_block:
            code_block_lines.append(line)
            continue

        # Handle markdown tables
        if stripped.startswith('|') and stripped.endswith('|'):
            # Table separator line e.g. | :--- | :--- |
            if re.match(r'^\|[\s\-:]+(\|[\s\-:]+)+\|$', stripped):
                continue
            cells = [c.strip() for c in stripped[1:-1].split('|')]
            table_rows.append(cells)
            continue
        else:
            if table_rows:
                # Render table
                num_rows = len(table_rows)
                num_cols = len(table_rows[0])
                table = doc.add_table(rows=num_rows, cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                
                # Column widths based on cols
                col_w = Inches(6.5 / num_cols)
                
                for r_idx, r_data in enumerate(table_rows):
                    row = table.rows[r_idx]
                    is_header = (r_idx == 0)
                    for c_idx, cell_value in enumerate(r_data):
                        cell = row.cells[c_idx]
                        cell.width = col_w
                        set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
                        set_cell_borders(cell, color="CBD5E1" if not is_header else "1B5E20")
                        
                        if is_header:
                            set_cell_background(cell, "1B5E20")
                        else:
                            bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                            set_cell_background(cell, bg)
                            
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.1
                        
                        # Handle <br> as newline
                        sub_lines = cell_value.replace('<br>', '\n').split('\n')
                        for sl_idx, sl in enumerate(sub_lines):
                            if sl_idx > 0:
                                p = cell.add_paragraph()
                                p.paragraph_format.space_before = Pt(1)
                                p.paragraph_format.space_after = Pt(1)
                                p.paragraph_format.line_spacing = 1.1
                            
                            run_tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', sl)
                            for t in run_tokens:
                                if not t:
                                    continue
                                if is_header:
                                    r = p.add_run(t.replace('**', ''))
                                    r.bold = True
                                    r.font.size = Pt(9.5)
                                    r.font.color.rgb = RGBColor(255, 255, 255)
                                else:
                                    if t.startswith('**') and t.endswith('**'):
                                        r = p.add_run(t[2:-2])
                                        r.bold = True
                                    elif t.startswith('*') and t.endswith('*'):
                                        r = p.add_run(t[1:-1])
                                        r.italic = True
                                    elif t.startswith('`') and t.endswith('`'):
                                        r = p.add_run(t[1:-1])
                                        r.font.name = 'Consolas'
                                        r.font.size = Pt(9)
                                        r.font.color.rgb = RGBColor(199, 37, 78)
                                    else:
                                        r = p.add_run(t)
                                    r.font.size = Pt(9.5)
                
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
                table_rows = []

        if not stripped:
            continue

        if stripped == '---':
            continue

        # Headings
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(27, 94, 32) # Dark Green

            if '5. Arsitektur Data' in h_text and os.path.exists('ERD_AgriBuddy.jpg'):
                # Add introductory text and embed ERD image
                intro_p = doc.add_paragraph()
                intro_p.paragraph_format.space_before = Pt(2)
                intro_p.paragraph_format.space_after = Pt(4)
                add_formatted_runs(intro_p, "Berikut adalah arsitektur model data konseptual dan relasional platform **AgriBuddy** yang menghubungkan 13 entitas utama ekosistem digital usahatani:")
                
                # Image container
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.paragraph_format.space_before = Pt(4)
                img_p.paragraph_format.space_after = Pt(2)
                img_p.add_run().add_picture('ERD_AgriBuddy.jpg', width=Inches(6.5))
                
                # Caption
                cap_p = doc.add_paragraph()
                cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_p.paragraph_format.space_before = Pt(0)
                cap_p.paragraph_format.space_after = Pt(8)
                cap_run = cap_p.add_run("Gambar 5.1: Entity Relationship Diagram (ERD) Platform AgriBuddy (13 Entitas)")
                cap_run.font.name = 'Calibri'
                cap_run.font.size = Pt(9)
                cap_run.italic = True
                cap_run.font.color.rgb = RGBColor(100, 100, 100)

            continue

        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(46, 125, 50)
            continue

        # Bullet list
        if stripped.startswith('* ') or stripped.startswith('- '):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            # Check indentation depth
            indent_level = (len(line) - len(line.lstrip())) // 2
            if indent_level > 0:
                p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            
            add_formatted_runs(p, item_text)
            continue

        # Numbered list
        m_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m_num:
            item_text = m_num.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, item_text)
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, stripped)

    # Save document
    output_filename = "PRD_AgriBuddy_Capstone.docx"
    doc.save(output_filename)
    print(f"Successfully generated: {output_filename} (Size: {os.path.getsize(output_filename)} bytes)")

    # Also copy to artifacts directory
    artifact_dir = r"C:\Users\Asus\.gemini\antigravity\brain\e07e4502-5d72-499c-989c-c0ab02ab60ae"
    dest_path = os.path.join(artifact_dir, output_filename)
    import shutil
    shutil.copy2(output_filename, dest_path)
    print(f"Copied to artifact directory: {dest_path}")

if __name__ == '__main__':
    build_docx()
